"""Synthèses mensuelles et export Excel consolidé."""
from datetime import date
from io import BytesIO

from django.db.models import Avg, Sum
from django.db.models.functions import TruncMonth

from cheptel.models import RapportJournalier
from finances import services as fin
from finances.models import Transaction
from stocks.models import MouvementStock

MOIS = ['Janvier', 'Février', 'Mars', 'Avril', 'Mai', 'Juin', 'Juillet', 'Août', 'Septembre', 'Octobre', 'Novembre', 'Décembre']


def synthese_mensuelle(ferme, months=12):
    """Une ligne par mois : production, mortalité, aliment, recettes, dépenses."""
    prod = {(r['m'].year, r['m'].month): r for r in
            RapportJournalier.objects.filter(ferme=ferme).annotate(m=TruncMonth('date')).values('m')
            .annotate(oeufs=Sum('oeufs_pondus'), morts=Sum('nombre_morts'), aliment=Avg('aliment_kg'))}
    series = fin.monthly_series(ferme, months=months)
    today = date.today()
    keys, y, m = [], today.year, today.month
    for _ in range(months):
        keys.append((y, m)); m -= 1
        if m == 0: y, m = y - 1, 12
    keys.reverse()
    rows = []
    for (y, m), s in zip(keys, series):
        p = prod.get((y, m), {})
        rows.append({'label': f'{MOIS[m - 1]} {y}', 'oeufs': p.get('oeufs') or 0, 'morts': p.get('morts') or 0,
                     'aliment': float(p.get('aliment') or 0), 'recettes': s['recettes'], 'depenses': s['depenses'],
                     'solde': s['recettes'] - s['depenses']})
    return rows[::-1]


def export_xlsx(ferme, with_finances=True):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    head_fill = PatternFill('solid', fgColor='16A34A')

    def sheet(ws, headers, rows, widths):
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF'); c.fill = head_fill; c.alignment = Alignment(horizontal='center')
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + i)].width = w
        ws.freeze_panes = 'A2'

    ws = wb.active; ws.title = 'Synthèse mensuelle'
    sheet(ws, ['Mois', 'Œufs', 'Morts', 'Aliment moyen (kg/j)', 'Recettes (FCFA)', 'Dépenses (FCFA)', 'Solde (FCFA)'],
          [[r['label'], r['oeufs'], r['morts'], round(r['aliment'], 1), r['recettes'], r['depenses'], r['solde']]
           for r in synthese_mensuelle(ferme)], [20, 12, 10, 22, 18, 18, 16])
    sheet(wb.create_sheet('Rapports journaliers'),
          ['Date', 'Œufs', 'Morts', 'Malades', 'Aliment (kg)', 'Eau (L)'],
          [[r.date, r.oeufs_pondus, r.nombre_morts, r.sujets_malades, r.aliment_kg, r.eau_litres]
           for r in RapportJournalier.objects.filter(ferme=ferme).order_by('-date')[:400]], [14, 10, 10, 10, 14, 12])
    sheet(wb.create_sheet('Stocks'), ['Date', 'Produit', 'Type', 'Quantité', 'Stock après', 'Motif'],
          [[m.date, m.produit.nom, m.get_type_display(), m.quantite, m.stock_apres, m.motif]
           for m in MouvementStock.objects.filter(produit__ferme=ferme).select_related('produit')[:500]],
          [14, 28, 12, 12, 14, 30])
    if with_finances:
      sheet(wb.create_sheet('Finances'), ['Date', 'Type', 'Catégorie', 'Libellé', 'Tiers', 'Montant (FCFA)'],
            [[t.date, t.get_type_display(), t.get_categorie_display(), t.libelle, t.tiers, t.montant]
             for t in Transaction.objects.filter(ferme=ferme).order_by('-date')[:800]], [14, 12, 24, 36, 26, 16])
    buf = BytesIO(); wb.save(buf)
    return buf.getvalue()
